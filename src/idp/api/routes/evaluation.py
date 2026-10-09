"""Evaluation suites (VRT-42) — measure the current models, prompts and
document types against documents whose right answer is known.
GET  /v1/eval-suites                    list, with each suite's latest run
GET  /v1/eval-suites/template           CSV header for a document type's fields
POST /v1/eval-suites                    from a CSV/Excel table + the files it names
POST /v1/eval-suites/from-corrections   golden set from the human corrections
GET  /v1/eval-suites/{id}               cases and runs
POST /v1/eval-suites/{id}/runs          run it (202), or resume its unfinished run
GET  /v1/eval-runs/{id}                 metrics and results; ?baseline= compares with another run"""

from __future__ import annotations

import csv
import io
import uuid
from datetime import UTC, datetime
from typing import Annotated, Any

from fastapi import APIRouter, BackgroundTasks, Depends, File, Form, HTTPException, Query, Response, UploadFile, status
from pydantic import BaseModel, Field
from sqlalchemy.ext.asyncio import AsyncSession

from idp.api.case_service import read_uploads
from idp.api.deps import get_app_settings, get_current_user, get_db_session, get_object_store, require_role
from idp.config import Settings
from idp.domain.evaluation import FILE_COLUMN, PAGES_COLUMN, TYPE_COLUMN, RunComparison, compare_runs, parse_table
from idp.evaluation.runner import is_active, outcome_of, run_evaluation
from idp.persistence.models import EvalCase, EvalRun, EvalSuite, User
from idp.persistence.repositories import DocumentTypeRepository, EvaluationRepository
from idp.storage.object_store import S3ObjectStore

router = APIRouter(prefix="/v1", tags=["evaluation"], dependencies=[Depends(get_current_user)])
_can_run = [Depends(require_role("operador", "admin"))]
_unprocessable = status.HTTP_422_UNPROCESSABLE_ENTITY


class RunSummary(BaseModel):
    id: uuid.UUID
    status: str
    created_by: str
    created_at: datetime
    finished_at: datetime | None
    error: str | None
    metrics: dict[str, Any] | None
    provenance: dict[str, Any] | None


class SuiteSummary(BaseModel):
    id: uuid.UUID
    name: str
    description: str | None
    source: str
    created_by: str
    created_at: datetime
    case_count: int
    latest_run: RunSummary | None


class CaseView(BaseModel):
    id: uuid.UUID
    filename: str
    page_start: int | None
    page_end: int | None
    expected_document_type: str | None
    expected_fields: dict[str, Any]
    source_document_id: uuid.UUID | None


class SuiteDetail(SuiteSummary):
    cases: list[CaseView]
    runs: list[RunSummary]


class ResultView(BaseModel):
    case_id: uuid.UUID
    filename: str
    expected_document_type: str | None
    predicted_document_type: str | None
    classification_confidence: float | None
    status: str
    error: str | None
    duration_ms: int | None
    fields: list[dict[str, Any]]


class RunDetail(RunSummary):
    suite_id: uuid.UUID
    suite_name: str
    case_count: int
    results: list[ResultView]
    baseline_id: uuid.UUID | None = None
    comparison: RunComparison | None = None


class GoldenSetRequest(BaseModel):
    name: str = Field(min_length=3)
    description: str | None = None


def _run_summary(run: EvalRun) -> RunSummary:
    return RunSummary(
        id=run.id, status=run.status, created_by=run.created_by, created_at=run.created_at, finished_at=run.finished_at,
        error=run.error, metrics=run.metrics, provenance=run.provenance,
    )


def _suite_summary(suite: EvalSuite) -> SuiteSummary:
    return SuiteSummary(
        id=suite.id, name=suite.name, description=suite.description, source=suite.source, created_by=suite.created_by,
        created_at=suite.created_at, case_count=len(suite.cases), latest_run=_run_summary(suite.runs[0]) if suite.runs else None,
    )


def _suite_detail(suite: EvalSuite) -> SuiteDetail:
    return SuiteDetail(
        **_suite_summary(suite).model_dump(),
        cases=[
            CaseView(
                id=c.id, filename=c.filename, page_start=c.page_start, page_end=c.page_end, expected_document_type=c.expected_document_type,
                expected_fields=c.expected_fields, source_document_id=c.source_document_id,
            )
            for c in suite.cases
        ],
        runs=[_run_summary(r) for r in suite.runs],
    )


def _read_table(table: UploadFile, content: bytes) -> list[dict[str, Any]]:
    name = (table.filename or "").lower()
    if name.endswith(".xlsx"):
        from openpyxl import load_workbook

        sheet = load_workbook(io.BytesIO(content), read_only=True, data_only=True).worksheets[0]
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return []
        header = [str(h).strip() if h is not None else "" for h in rows[0]]
        return [{h: v for h, v in zip(header, row, strict=False) if h} for row in rows[1:]]
    if name.endswith(".csv"):
        text = content.decode("utf-8-sig")
        dialect = csv.Sniffer().sniff(text.splitlines()[0] if text else ",", delimiters=",;")
        return list(csv.DictReader(io.StringIO(text), dialect=dialect))
    raise HTTPException(status_code=_unprocessable, detail="la tabla debe ser .csv o .xlsx")


async def _suite_or_404(session: AsyncSession, suite_id: uuid.UUID) -> EvalSuite:
    suite = await EvaluationRepository(session).get_suite(suite_id)
    if suite is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="suite not found")
    return suite


@router.get("/eval-suites", response_model=list[SuiteSummary])
async def list_suites(session: AsyncSession = Depends(get_db_session)) -> list[SuiteSummary]:
    return [_suite_summary(s) for s in await EvaluationRepository(session).list_suites()]


@router.get("/eval-suites/template")
async def table_template(document_type: str, session: AsyncSession = Depends(get_db_session)) -> Response:
    current = (await DocumentTypeRepository(session).load_catalog()).current(document_type)
    if current is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail=f"el tipo '{document_type}' no está en el catálogo")
    header = [FILE_COLUMN, TYPE_COLUMN, PAGES_COLUMN, *(f.name for f in current[1].fields if f.type != "list")]
    return Response(
        ",".join(header) + "\n",
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": f'attachment; filename="evaluacion-{document_type}.csv"'},
    )


@router.post("/eval-suites", response_model=SuiteDetail, status_code=status.HTTP_201_CREATED, dependencies=_can_run)
async def create_suite(
    name: Annotated[str, Form(min_length=3)],
    table: Annotated[UploadFile, File(description="CSV o Excel: archivo, tipo, paginas y una columna por campo.")],
    files: Annotated[list[UploadFile], File(description="Los documentos que nombra la columna 'archivo'.")],
    description: Annotated[str | None, Form()] = None,
    session: AsyncSession = Depends(get_db_session),
    object_store: S3ObjectStore = Depends(get_object_store),
    user: User = Depends(get_current_user),
) -> SuiteDetail:
    specs, errors = parse_table(_read_table(table, await table.read()), await DocumentTypeRepository(session).load_catalog())
    uploads = {u.filename: u for u in await read_uploads(files)}
    errors += [f"'{name}' no está entre los archivos subidos" for name in sorted({s.filename for s in specs} - set(uploads))]
    if errors:
        raise HTTPException(status_code=_unprocessable, detail="; ".join(errors))

    suite_id = uuid.uuid4()
    keys: dict[str, str] = {}
    for filename in {s.filename for s in specs}:
        upload = uploads[filename]
        keys[filename] = f"default/eval/{suite_id}/{uuid.uuid4()}/{filename}"
        object_store.put(keys[filename], upload.content, content_type=upload.content_type)
    cases = [
        EvalCase(
            position=i, filename=s.filename, storage_key=keys[s.filename], page_start=s.page_start, page_end=s.page_end,
            expected_document_type=s.expected_document_type, expected_fields=s.expected_fields,
        )
        for i, s in enumerate(specs)
    ]
    suite = await EvaluationRepository(session).create_suite(
        name=name, description=description, source="table", created_by=user.name, cases=cases, suite_id=suite_id
    )
    await session.commit()
    return _suite_detail(await _suite_or_404(session, suite.id))


@router.post("/eval-suites/from-corrections", response_model=SuiteDetail, status_code=status.HTTP_201_CREATED, dependencies=_can_run)
async def create_golden_set(
    body: GoldenSetRequest, session: AsyncSession = Depends(get_db_session), user: User = Depends(get_current_user)
) -> SuiteDetail:
    """Each document a human corrected, expected to produce what the human
    left: its type and the corrected fields."""
    repo = EvaluationRepository(session)
    corrected = await repo.corrected_documents()
    if not corrected:
        raise HTTPException(status_code=_unprocessable, detail="todavía no hay correcciones con las que armar un golden set")
    cases = [
        EvalCase(
            position=i, filename=d.original_filename, storage_key=d.storage_key, page_start=d.page_start, page_end=d.page_end,
            expected_document_type=d.document_type, expected_fields=fields, source_document_id=d.id,
        )
        for i, (d, fields) in enumerate(corrected)
    ]
    description = body.description or f"{len(cases)} documentos corregidos hasta {datetime.now(UTC):%Y-%m-%d %H:%M} UTC"
    suite = await repo.create_suite(name=body.name, description=description, source="corrections", created_by=user.name, cases=cases)
    await session.commit()
    return _suite_detail(await _suite_or_404(session, suite.id))


@router.get("/eval-suites/{suite_id}", response_model=SuiteDetail)
async def get_suite(suite_id: uuid.UUID, session: AsyncSession = Depends(get_db_session)) -> SuiteDetail:
    return _suite_detail(await _suite_or_404(session, suite_id))


@router.post("/eval-suites/{suite_id}/runs", response_model=RunSummary, status_code=status.HTTP_202_ACCEPTED, dependencies=_can_run)
async def start_run(
    suite_id: uuid.UUID,
    background: BackgroundTasks,
    session: AsyncSession = Depends(get_db_session),
    settings: Settings = Depends(get_app_settings),
    user: User = Depends(get_current_user),
) -> RunSummary:
    suite = await _suite_or_404(session, suite_id)
    repo = EvaluationRepository(session)
    run = await repo.unfinished_run(suite.id)
    if run is None:
        run = EvalRun(suite_id=suite.id, created_by=user.name)
        session.add(run)
        await session.commit()
        await session.refresh(run)
    if not is_active(run.id):
        background.add_task(run_evaluation, settings, run.id)
    return _run_summary(run)


@router.get("/eval-runs/{run_id}", response_model=RunDetail)
async def get_run(
    run_id: uuid.UUID,
    baseline: Annotated[uuid.UUID | None, Query(description="Otra corrida de la misma suite con la cual comparar.")] = None,
    session: AsyncSession = Depends(get_db_session),
) -> RunDetail:
    repo = EvaluationRepository(session)
    run = await repo.get_run(run_id)
    if run is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="run not found")
    cases = {c.id: c for c in run.suite.cases}
    detail = RunDetail(
        **_run_summary(run).model_dump(),
        suite_id=run.suite_id,
        suite_name=run.suite.name,
        case_count=len(cases),
        results=[
            ResultView(
                case_id=r.case_id, filename=cases[r.case_id].filename, expected_document_type=cases[r.case_id].expected_document_type,
                predicted_document_type=r.predicted_document_type, classification_confidence=r.classification_confidence,
                status=r.status, error=r.error, duration_ms=r.duration_ms, fields=r.field_results,
            )
            for r in sorted(run.results, key=lambda r: cases[r.case_id].position)
        ],
    )
    if baseline is not None:
        base = await repo.get_run(baseline)
        if base is None or base.suite_id != run.suite_id:
            raise HTTPException(status_code=_unprocessable, detail="la corrida base debe ser de la misma suite")
        detail.baseline_id = base.id
        detail.comparison = compare_runs(
            [outcome_of(r, cases[r.case_id]) for r in base.results], [outcome_of(r, cases[r.case_id]) for r in run.results]
        )
    return detail
