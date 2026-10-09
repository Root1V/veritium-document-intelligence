"""Tables people fill in Excel: a .csv (comma or semicolon) or a .xlsx's
first sheet, read as one dict per row keyed by the header. Shared by
evaluation suites (VRT-42) and bulk manifests (VRT-48)."""

from __future__ import annotations

import csv
import io
from typing import Any


def read_table(filename: str, content: bytes) -> list[dict[str, Any]]:
    """Raises ValueError for any other format."""
    name = filename.lower()
    if name.endswith(".xlsx"):
        from openpyxl import load_workbook

        sheet = load_workbook(io.BytesIO(content), read_only=True, data_only=True).worksheets[0]
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return []
        header = [str(h).strip() if h is not None else "" for h in rows[0]]
        return [{h: v for h, v in zip(header, row, strict=False) if h} for row in rows[1:]]
    if name.endswith(".csv"):
        text = content.decode("utf-8-sig")
        dialect = csv.Sniffer().sniff(text.splitlines()[0] if text else ",", delimiters=",;")
        return list(csv.DictReader(io.StringIO(text), dialect=dialect))
    raise ValueError("la tabla debe ser .csv o .xlsx")
